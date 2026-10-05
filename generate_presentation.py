import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6] # Blank slide

    # Brand Colors
    BG_COLOR = RGBColor(11, 15, 25)         # Deep slate/black
    CARD_BG = RGBColor(19, 26, 44)          # Elevated card navy
    BORDER_COLOR = RGBColor(38, 50, 80)     # Subtle border
    ACCENT_INDIGO = RGBColor(99, 102, 241)  # #6366f1 Brand Indigo
    ACCENT_VIOLET = RGBColor(139, 92, 246)  # #8b5cf6 Violet
    ACCENT_CYAN = RGBColor(6, 182, 212)     # #06b6d4 Cyan
    ACCENT_EMERALD = RGBColor(16, 185, 129) # #10b981 Emerald
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184)    # Slate 400
    TEXT_LIGHT = RGBColor(226, 232, 240)    # Slate 200

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.color.rgb = BG_COLOR
        return bg

    def add_header(slide, tag_text, title_text, category_color=ACCENT_INDIGO):
        # Badge / Tag
        tag_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = category_color

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.9), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

    # ----------------------------------------------------
    # SLIDE 1: Title & Problem Statement
    # ----------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    # Big Pill Tag
    pill = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.1), Inches(4.2), Inches(0.45))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(29, 38, 64)
    pill.line.color.rgb = ACCENT_INDIGO
    pill_tf = pill.text_frame
    pill_p = pill_tf.paragraphs[0]
    pill_p.text = "AI-POWERED TECHNICAL INTERVIEW AGENT"
    pill_p.font.name = "Segoe UI"
    pill_p.font.size = Pt(11)
    pill_p.font.bold = True
    pill_p.font.color.rgb = ACCENT_CYAN
    pill_p.alignment = PP_ALIGN.CENTER

    # Title Text
    t_box = s1.shapes.add_textbox(Inches(0.9), Inches(1.75), Inches(11.5), Inches(1.6))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "INTERVIEWPILOT AI"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(46)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Autonomous, Adaptive Technical Interview Agent for College Placements"
    p1_sub.font.name = "Segoe UI"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = ACCENT_INDIGO
    p1_sub.space_before = Pt(8)

    # 2 Comparison Cards: Problem vs Solution
    # Left Card: Traditional Mock Tests
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(3.7), Inches(5.5), Inches(2.9))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = BORDER_COLOR
    c1_tf = card1.text_frame
    c1_tf.word_wrap = True
    c1_p0 = c1_tf.paragraphs[0]
    c1_p0.text = "THE PROBLEM: Static Mock Tests"
    c1_p0.font.bold = True
    c1_p0.font.size = Pt(15)
    c1_p0.font.color.rgb = RGBColor(244, 63, 94) # Coral Red
    
    bullets_problem = [
        "Fixed Question Banks: Repeated, memorized questions with no variation.",
        "Zero Real-Time Adaptation: Difficulty stays identical regardless of performance.",
        "Superficial Scoring: Multiple-choice quizzes miss technical depth and reasoning.",
        "Generic Feedback: Students don't receive targeted diagnostic roadmaps."
    ]
    for b in bullets_problem:
        p = c1_tf.add_paragraph()
        p.text = f"- {b}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(6)

    # Right Card: InterviewPilot AI Solution
    card2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.7), Inches(5.6), Inches(2.9))
    card2.fill.solid()
    card2.fill.fore_color.rgb = CARD_BG
    card2.line.color.rgb = ACCENT_INDIGO
    c2_tf = card2.text_frame
    c2_tf.word_wrap = True
    c2_p0 = c2_tf.paragraphs[0]
    c2_p0.text = "THE SOLUTION: Autonomous Agentic Interviewer"
    c2_p0.font.bold = True
    c2_p0.font.size = Pt(15)
    c2_p0.font.color.rgb = ACCENT_EMERALD

    bullets_solution = [
        "LLM-Powered Orchestration: Powered by Google Gemini 1.5 Flash.",
        "Dynamic Difficulty Adjustment: Automatically scales beginner to advanced.",
        "Multi-Dimensional Rubric: Evaluates depth, accuracy, structure, and clarity.",
        "Personalized Diagnostic Dossier: 100-pt benchmark & 4-milestone study plan."
    ]
    for b in bullets_solution:
        p = c2_tf.add_paragraph()
        p.text = f"- {b}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_before = Pt(6)

    # Speaker notes for Slide 1
    s1.notes_slide.notes_text_frame.text = (
        "Good morning respected faculty. Today I present InterviewPilot AI, an autonomous, "
        "LLM-powered technical interview agent built to prepare engineering and MCA students for real technical interviews."
    )

    # ----------------------------------------------------
    # SLIDE 2: Closed-Loop Agentic Workflow
    # ----------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "Agent Architecture", "The Closed-Loop Feedback Workflow")

    steps = [
        ("01. SETUP & CALIBRATION", "Candidate specifies target role (e.g. SDE Fresher), level, and technical skills (Python, DSA, DBMS, OS).", ACCENT_CYAN),
        ("02. QUESTION GENERATION", "Agent synthesizes a realistic, role-specific question aligned with current target difficulty.", ACCENT_INDIGO),
        ("03. REAL-TIME EVALUATION", "LLM analyzes response against 4-pillar rubric, calculating score (0-10) and feedback.", ACCENT_VIOLET),
        ("04. ADAPTIVE STATE SHIFT", "Dynamic Difficulty Engine escalates or de-escalates next question and flags weak concepts.", ACCENT_EMERALD),
    ]

    left_start = Inches(0.9)
    card_width = Inches(2.7)
    gap = Inches(0.24)

    for i, (stitle, sdesc, scolor) in enumerate(steps):
        c_left = left_start + (card_width + gap) * i
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, Inches(2.1), card_width, Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = scolor
        card.line.width = Pt(1.5)

        tf = card.text_frame
        tf.word_wrap = True
        
        # Step header
        p0 = tf.paragraphs[0]
        p0.text = stitle
        p0.font.bold = True
        p0.font.size = Pt(13)
        p0.font.color.rgb = scolor

        # Step description
        p1 = tf.add_paragraph()
        p1.text = sdesc
        p1.font.size = Pt(12)
        p1.font.color.rgb = TEXT_LIGHT
        p1.space_before = Pt(16)

        # Micro indicator at bottom
        p2 = tf.add_paragraph()
        p2.text = f"Phase {i+1} of 4"
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(40)

    s2.notes_slide.notes_text_frame.text = (
        "Unlike static question forms, this system operates on a closed-loop feedback architecture. "
        "It initiates an interview calibrated to the student's role, evaluates answers in real-time, and dynamically steers subsequent questions."
    )

    # ----------------------------------------------------
    # SLIDE 3: 4-Pillar Rubric & Dynamic Difficulty (DDA)
    # ----------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Evaluation & Adaptation", "4-Pillar Rubric & Dynamic Difficulty Adjustment")

    # Left Column: Rubric
    r_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.0), Inches(5.6), Inches(4.7))
    r_card.fill.solid()
    r_card.fill.fore_color.rgb = CARD_BG
    r_card.line.color.rgb = BORDER_COLOR
    r_tf = r_card.text_frame
    r_tf.word_wrap = True
    
    rp0 = r_tf.paragraphs[0]
    rp0.text = "THE 4-PILLAR SCORING RUBRIC"
    rp0.font.bold = True
    rp0.font.size = Pt(16)
    rp0.font.color.rgb = ACCENT_INDIGO

    rubrics = [
        ("1. Technical Depth (30%)", "Evaluates understanding of runtime complexity, memory layout, system design, and edge cases."),
        ("2. Correctness & Accuracy (30%)", "Checks factual syntax, algorithm correctness, and proper computer science principles."),
        ("3. Communication & Structure (20%)", "Rewards structured articulation, logical steps, and clear reasoning."),
        ("4. Relevance & Intent (20%)", "Ensures the candidate directly answers the question without dodging key points.")
    ]
    for r_title, r_desc in rubrics:
        p_t = r_tf.add_paragraph()
        p_t.text = r_title
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(10)

        p_d = r_tf.add_paragraph()
        p_d.text = r_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    # Right Column: Dynamic Difficulty Adjustment
    d_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.6), Inches(4.7))
    d_card.fill.solid()
    d_card.fill.fore_color.rgb = CARD_BG
    d_card.line.color.rgb = ACCENT_EMERALD
    d_tf = d_card.text_frame
    d_tf.word_wrap = True

    dp0 = d_tf.paragraphs[0]
    dp0.text = "DYNAMIC DIFFICULTY ADJUSTMENT (DDA)"
    dp0.font.bold = True
    dp0.font.size = Pt(16)
    dp0.font.color.rgb = ACCENT_EMERALD

    dda_tiers = [
        ("Score >= 8.0 / 10 -> ESCALATE DIFFICULTY", "Beginner -> Intermediate -> Advanced. Probes complex edge cases, concurrency, and trade-offs.", ACCENT_EMERALD),
        ("Score 5.0 - 7.9 / 10 -> MAINTAIN TIER", "Keeps current difficulty to verify domain consistency and stability across diverse topics.", ACCENT_CYAN),
        ("Score < 5.0 / 10 -> DE-ESCALATE & DIAGNOSE", "Reduces complexity to core fundamentals; flags concept as a 'Weak Skill' for targeted follow-ups.", RGBColor(251, 146, 60))
    ]
    for d_title, d_desc, col in dda_tiers:
        p_t = d_tf.add_paragraph()
        p_t.text = d_title
        p_t.font.bold = True
        p_t.font.size = Pt(13)
        p_t.font.color.rgb = col
        p_t.space_before = Pt(14)

        p_d = d_tf.add_paragraph()
        p_d.text = d_desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    s3.notes_slide.notes_text_frame.text = (
        "The evaluation engine scores answers across four dimensions: Technical Depth, Correctness, Communication, and Relevance. "
        "If a candidate scores 8 or above, difficulty escalates. If they score below 5, the agent dials back difficulty to probe core fundamentals."
    )

    # ----------------------------------------------------
    # SLIDE 4: Context Memory & Weakness Targeted Probing
    # ----------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "Intelligent Memory", "Context Memory, Zero Repeats & Weak-Spot Isolation")

    features = [
        ("ZERO REPETITION GUARANTEE", "Session history stores asked questions and specific subtopics. The prompt engine strictly bans duplicate topics, guaranteeing broad placement coverage.", ACCENT_INDIGO),
        ("WEAK SPOT ISOLATION", "When a low score is detected in a skill (e.g. SQL indexing or Dynamic Programming), the agent actively steers upcoming turns back to that weak domain to reinforce it.", ACCENT_CYAN),
        ("VOICE-FIRST MULTI-MODAL INTERACTION", "Equipped with real-time Speech Synthesis (TTS) for conversational AI voice narration and Web Speech Recognition for hands-free verbal answering.", ACCENT_VIOLET)
    ]

    for i, (ftitle, fdesc, fcol) in enumerate(features):
        top_y = Inches(2.0 + i * 1.55)
        f_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), top_y, Inches(11.5), Inches(1.35))
        f_card.fill.solid()
        f_card.fill.fore_color.rgb = CARD_BG
        f_card.line.color.rgb = fcol
        
        f_tf = f_card.text_frame
        f_tf.word_wrap = True

        fp0 = f_tf.paragraphs[0]
        fp0.text = ftitle
        fp0.font.bold = True
        fp0.font.size = Pt(14)
        fp0.font.color.rgb = fcol

        fp1 = f_tf.add_paragraph()
        fp1.text = fdesc
        fp1.font.size = Pt(12)
        fp1.font.color.rgb = TEXT_LIGHT
        fp1.space_before = Pt(4)

    s4.notes_slide.notes_text_frame.text = (
        "Through session memory, the agent guarantees zero repetitive questions and isolates weak topic areas "
        "to challenge the student where they struggle most, complete with voice narration."
    )

    # ----------------------------------------------------
    # SLIDE 5: Architecture & Hiring Dossier
    # ----------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "System & Outcome", "Production Stack & The Final Hiring Dossier")

    # Left: Production Tech Stack
    t_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.0), Inches(5.6), Inches(4.7))
    t_card.fill.solid()
    t_card.fill.fore_color.rgb = CARD_BG
    t_card.line.color.rgb = ACCENT_INDIGO
    t_tf = t_card.text_frame
    t_tf.word_wrap = True

    tp0 = t_tf.paragraphs[0]
    tp0.text = "FULL-STACK PRODUCTION ARCHITECTURE"
    tp0.font.bold = True
    tp0.font.size = Pt(16)
    tp0.font.color.rgb = ACCENT_INDIGO

    stack_items = [
        ("Frontend Application", "React 18, TypeScript, Tailwind CSS, Lucide icons, Vite 6 build system."),
        ("AI Intelligence Engine", "Google Gemini 1.5 Flash structured JSON model."),
        ("Resilient Fallback Engine", "Autonomous offline heuristics for 100% test uptime with zero downtime."),
        ("Serverless Cloud Deployment", "Vercel Serverless Functions with zero client-side API key exposure.")
    ]
    for s_title, s_desc in stack_items:
        p_t = t_tf.add_paragraph()
        p_t.text = f"- {s_title}:"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(10)

        p_d = t_tf.add_paragraph()
        p_d.text = f"  {s_desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    # Right: Comprehensive Final Report Deliverables
    o_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.0), Inches(5.6), Inches(4.7))
    o_card.fill.solid()
    o_card.fill.fore_color.rgb = CARD_BG
    o_card.line.color.rgb = ACCENT_EMERALD
    o_tf = o_card.text_frame
    o_tf.word_wrap = True

    op0 = o_tf.paragraphs[0]
    op0.text = "ACTIONABLE HIRING REPORT DELIVERABLES"
    op0.font.bold = True
    op0.font.size = Pt(16)
    op0.font.color.rgb = ACCENT_EMERALD

    report_items = [
        ("100-Point Score & Recommendation", "Strong Hire (>=85), Hire (>=70), Borderline (>=50), Needs Improvement."),
        ("Multi-Metric Breakdown", "Separates Technical Depth, Communication, and Problem-Solving scores."),
        ("Question-Level Audit", "Displays every answer, evaluation breakdown, strengths, and missed points."),
        ("4-Milestone Study Roadmap", "High & Medium priority technical revision plan tailored to candidate gaps.")
    ]
    for r_title, r_desc in report_items:
        p_t = o_tf.add_paragraph()
        p_t.text = f"- {r_title}:"
        p_t.font.bold = True
        p_t.font.size = Pt(12)
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(10)

        p_d = o_tf.add_paragraph()
        p_d.text = f"  {r_desc}"
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_LIGHT
        p_d.space_before = Pt(2)

    s5.notes_slide.notes_text_frame.text = (
        "Built with React, TypeScript, and Google Gemini on a secure serverless backend, "
        "it concludes by generating a hiring dossier with a 100-point benchmark and a targeted study roadmap. Thank you!"
    )

    # Save outputs
    out_root = "InterviewPilot_AI_Presentation.pptx"
    os.makedirs("public", exist_ok=True)
    out_public = os.path.join("public", "InterviewPilot_AI_Presentation.pptx")
    prs.save(out_root)
    prs.save(out_public)
    print(f"Presentation successfully created at {out_root} and {out_public}")

if __name__ == "__main__":
    create_deck()
